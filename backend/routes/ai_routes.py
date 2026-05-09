"""
routes/ai_routes.py - AI-powered career endpoints (resume, cover letter, interview prep, career paths).

Endpoints (under /api):
  POST   /ai/optimize-resume               -- Tailor resume to a job description
  POST   /ai/detailed-match-analysis       -- Resume-grounded match analysis
  POST   /ai/cover-letter                  -- Generate a cover letter
  POST   /ai/download-docx                 -- Convert text to .docx for download
  POST   /ai/interview-prep                -- Generate interview prep guide
  POST   /ai/career-paths                  -- Analyze and suggest career paths
  DELETE /ai/career-paths                  -- Clear career path analysis
  GET    /ai/learning-resources            -- Suggested learning resources
  GET    /ai/career-paths/{path_title}/jobs-- Real job listings for a career path

Refactored from monolithic server.py (Feb 2026).
"""
from datetime import datetime, timezone
from typing import List, Optional, Dict, Any
import asyncio
import io
import json
import re
import uuid

from fastapi import APIRouter, HTTPException, Request
from fastapi.responses import StreamingResponse
import httpx

from core import db, logger, EMERGENT_LLM_KEY, RAPIDAPI_KEY, get_current_user
from models import (
    OptimizeResumeRequest,
    DetailedMatchRequest,
    GenerateCoverLetterRequest,
    DocxDownloadRequest,
    InterviewPrepRequest,
)
from profile_schema import get_skill_names
from encryption import decrypt_sensitive_data
from stripe_routes import check_usage_limit, increment_usage
from routes.job_routes import evaluate_job_match

router = APIRouter()


@router.post("/ai/optimize-resume")
async def optimize_resume(request: Request, req: OptimizeResumeRequest):
    """Optimize resume for ATS based on job description while preserving original format."""
    user = await get_current_user(request)
    
    # Check usage limits
    usage_check = await check_usage_limit(user.user_id, "resume_optimizations")
    if not usage_check["allowed"]:
        raise HTTPException(status_code=402, detail={
            "error": "usage_limit_reached",
            "feature": "resume_optimizations",
            "current": usage_check["current"],
            "limit": usage_check["limit"],
            "message": f"You've used all {usage_check['limit']} resume optimizations this month. Upgrade to Pro for unlimited access."
        })
    
    profile = await db.user_profiles.find_one(
        {"user_id": user.user_id},
        {"_id": 0}
    )
    
    if not profile or not profile.get("resume_text"):
        raise HTTPException(status_code=400, detail="Please upload your resume first")
    
    from emergentintegrations.llm.chat import LlmChat, UserMessage
    
    resume_format = profile.get("resume_format", "text")
    original_resume = profile.get("resume_text", "")
    
    chat = LlmChat(
        api_key=EMERGENT_LLM_KEY,
        session_id=f"resume_opt_{user.user_id}_{uuid.uuid4().hex[:8]}",
        system_message="""You are an expert ATS (Applicant Tracking System) resume optimizer.

**CRITICAL: ONE PAGE MAXIMUM**
The final resume MUST fit on ONE SINGLE PAGE when pasted into a Word document with standard margins (1 inch) and 11pt font. This is NON-NEGOTIABLE. Be concise and prioritize the most impactful content.

PAGE LENGTH GUIDELINES:
- Maximum ~450-500 words total
- 3-4 bullet points per role (not 5-6)
- Keep bullet points to 1-2 lines max
- Trim older/less relevant experience if needed to fit
- Summary/objective should be 2-3 lines max

FORMATTING RULES:
1. Preserve the original resume's structure and section order
2. Keep section headers in the same format
3. Use bullet points (•) or dashes (-) consistently
4. Maintain clear spacing between sections
5. Keep date formats consistent (e.g., "Jan 2020 - Present")
6. Do NOT add new sections or reorganize

OPTIMIZATION FOCUS:
- Inject relevant keywords from the job description naturally
- Strengthen action verbs
- Include quantifiable metrics where possible
- Mirror terminology from the job description
- PRIORITIZE recent and most relevant experience
- CUT less impactful content to fit one page

OUTPUT FORMAT:
Return the optimized resume as CONCISE plain text that fits on ONE PAGE."""
    ).with_model("openai", "gpt-5.2")
    
    prompt = f"""Optimize this resume for the following job. The output MUST FIT ON ONE SINGLE PAGE in a Word document.

JOB DESCRIPTION:
{req.job_description}

ORIGINAL RESUME:
{original_resume}

CRITICAL REQUIREMENTS:
1. **ONE PAGE ONLY** - This is the most important rule. Be concise!
2. Keep section headers and overall structure
3. Use consistent bullet points (• or -)
4. Limit each role to 3-4 impactful bullet points (1-2 lines each)
5. Enhance content with keywords from the job description
6. Cut less relevant content if needed to fit one page
7. Maximum ~450-500 words total

Return the concise, one-page optimized resume now:"""
    
    try:
        response = await chat.send_message(UserMessage(text=prompt))
        await increment_usage(user.user_id, "resume_optimizations")
        return {"optimized_resume": response, "original_format": resume_format}
    except Exception as e:
        logger.error(f"Resume optimization error: {str(e)}")
        raise HTTPException(status_code=500, detail="Failed to optimize resume")



@router.post("/ai/detailed-match-analysis")
async def generate_detailed_match_analysis(request: Request, req: DetailedMatchRequest):
    """
    Generate resume-grounded match analysis.
    Every strength must reference specific resume evidence and job requirements.
    """
    user = await get_current_user(request)
    
    profile = await db.user_profiles.find_one(
        {"user_id": user.user_id},
        {"_id": 0}
    )
    
    if not profile:
        raise HTTPException(status_code=400, detail="Please complete your profile first")
    
    resume_text = profile.get("resume_text", "")
    if not resume_text or len(resume_text) < 100:
        raise HTTPException(status_code=400, detail="Please upload your resume for detailed analysis")
    
    # Get skills with years for context
    skills = profile.get("skills", [])
    skills_text = ""
    if skills:
        if isinstance(skills[0], dict):
            skills_text = ", ".join([f"{s.get('name', '')} ({s.get('years', 0)} years)" for s in skills])
        else:
            skills_text = ", ".join(skills)
    
    experience_years = profile.get("experience_years", 0)
    target_roles = profile.get("job_titles", [])
    seniority = profile.get("seniority_level", "mid")
    
    from emergentintegrations.llm.chat import LlmChat, UserMessage
    
    chat = LlmChat(
        api_key=EMERGENT_LLM_KEY,
        session_id=f"match_analysis_{user.user_id}_{uuid.uuid4().hex[:8]}",
        system_message="""You are an expert career analyst who provides evidence-based job match analysis.

CRITICAL RULES:
1. Every strength MUST reference specific evidence from the resume
2. Every strength MUST reference a specific job requirement
3. NO generic phrases like "Strong role alignment", "Experience aligns well", "Good fit"
4. If resume evidence is weak or missing for a requirement, say so explicitly
5. Do NOT fabricate or assume experience not in the resume
6. Prefer fewer, higher-quality bullets over generic coverage

FORMAT for each strength bullet:
[Job requirement] → [Resume evidence] → [Why it matters]

Example of GOOD analysis:
"SQL-based analysis required → Resume shows 'Built SQL queries for performance reporting and trend analysis' → Direct skill match for analytical requirements"

Example of BAD analysis (DO NOT DO THIS):
"Strong role alignment: Your experience aligns well with this position" 
"""
    ).with_model("openai", "gpt-5.2")
    
    prompt = f"""Analyze this job match based on the candidate's ACTUAL resume content.

JOB DETAILS:
- Title: {req.job_title}
- Company: {req.company_name}
- Description: {req.job_description}

CANDIDATE PROFILE:
- Experience: {experience_years} years
- Seniority Level: {seniority}
- Target Roles: {', '.join(target_roles) if target_roles else 'Not specified'}
- Skills: {skills_text if skills_text else 'Not specified'}

CANDIDATE RESUME (use this as evidence):
{resume_text}

GENERATE A RESUME-GROUNDED MATCH ANALYSIS:

1. **Match Summary** (1-2 sentences)
   Explain why this role fits based on SPECIFIC resume evidence, not generic statements.

2. **Strengths** (3-5 bullets, ONLY if evidence exists)
   Each bullet MUST follow this format:
   [Specific job requirement] → [Specific resume bullet/experience] → [Why this matters for the role]
   
   Focus on:
   - Tools/technologies mentioned in both job and resume
   - Metrics/outcomes from resume that match job needs
   - Domain experience that aligns
   - Relevant certifications or education

3. **Gaps/Concerns** (1-3 bullets, be honest)
   List requirements from the job description where:
   - Resume evidence is weak or missing
   - Experience level may not match
   - Skills need development

4. **Recommendation**
   Should the candidate apply? Why or why not based on evidence?

IMPORTANT: Be specific and honest. Reference actual text from the resume. If you can't find evidence for something, say "Resume does not demonstrate..." rather than making assumptions."""

    try:
        response = await chat.send_message(UserMessage(text=prompt))
        return {
            "detailed_analysis": response,
            "job_title": req.job_title,
            "company": req.company_name,
            "skills_analyzed": skills_text,
            "experience_years": experience_years
        }
    except Exception as e:
        logger.error(f"Detailed match analysis error: {str(e)}")
        raise HTTPException(status_code=500, detail="Failed to generate detailed analysis")


@router.post("/ai/cover-letter")
async def generate_cover_letter(request: Request, req: GenerateCoverLetterRequest):
    """Generate personalized cover letter."""
    user = await get_current_user(request)
    
    # Check usage limits
    usage_check = await check_usage_limit(user.user_id, "cover_letters")
    if not usage_check["allowed"]:
        raise HTTPException(status_code=402, detail={
            "error": "usage_limit_reached",
            "feature": "cover_letters",
            "current": usage_check["current"],
            "limit": usage_check["limit"],
            "message": f"You've used all {usage_check['limit']} cover letter generations this month. Upgrade to Pro for unlimited access."
        })
    
    profile = await db.user_profiles.find_one(
        {"user_id": user.user_id},
        {"_id": 0}
    )
    
    user_doc = await db.users.find_one(
        {"user_id": user.user_id},
        {"_id": 0}
    )
    
    from emergentintegrations.llm.chat import LlmChat, UserMessage
    
    chat = LlmChat(
        api_key=EMERGENT_LLM_KEY,
        session_id=f"cover_{user.user_id}_{uuid.uuid4().hex[:8]}",
        system_message="""You are an expert ATS-optimized cover letter writer. Write concise, professional cover letters that directly align candidate experience to job requirements.

STRICT RULES:
- 1 page max (300-450 words)
- Simple formatting (no tables, no columns, no emojis)
- No fluff or generic enthusiasm
- Use keywords directly from the job description
- Align experience clearly to role requirements
- Professional, confident tone (not desperate or salesy)
- No company clichés or buzzwords
- AVOID phrases like "I am passionate", "I am excited", "I would love to", "I am thrilled"

REQUIRED STRUCTURE:
1. Opening paragraph: State the role title and company. Briefly summarize why the candidate's background fits the role.
2. Middle paragraphs (1-2): Match experience directly to key job requirements. Use metrics or outcomes where possible. Mirror terminology used in the job description.
3. Closing paragraph: Reiterate fit. Express interest in discussing the role. Thank the reader.

Output only the cover letter text, no additional commentary."""
    ).with_model("openai", "gpt-5.2")
    
    # Handle v2 skills format (list of dicts with name, years, level)
    raw_skills = profile.get("skills", []) if profile else []
    skill_names = get_skill_names(raw_skills)
    skills = ", ".join(skill_names) if skill_names else "Not specified"
    experience = profile.get("experience_years", 0) if profile else 0
    resume = profile.get("resume_text", "") if profile else ""
    job_titles = ", ".join(profile.get("job_titles", [])) if profile else "Not specified"
    
    prompt = f"""Write an ATS-friendly cover letter using these inputs:

JOB TITLE: {req.job_title}
COMPANY: {req.company}

JOB DESCRIPTION:
{req.job_description}

CANDIDATE INFORMATION:
- Name: {user_doc['name']}
- Years of Experience: {experience}
- Key Skills: {skills}
- Target Roles: {job_titles}

RESUME CONTENT:
{resume[:2000] if resume else 'Not provided'}

Generate a professional, ATS-optimized cover letter following the strict rules and structure provided. Use keywords from the job description and align the candidate's experience directly to the role requirements."""
    
    try:
        response = await chat.send_message(UserMessage(text=prompt))
        await increment_usage(user.user_id, "cover_letters")
        return {"cover_letter": response}
    except Exception as e:
        logger.error(f"Cover letter generation error: {str(e)}")
        raise HTTPException(status_code=500, detail="Failed to generate cover letter")



@router.post("/ai/download-docx")
async def download_docx(request: Request, req: DocxDownloadRequest):
    """Convert AI-generated text to a formatted .docx file."""
    await get_current_user(request)
    
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    import io
    
    doc = Document()
    
    # Set default margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)
    
    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(33, 33, 33)
    
    style.paragraph_format.space_after = Pt(2)
    style.paragraph_format.space_before = Pt(0)
    
    lines = req.content.strip().split('\n')
    
    if req.doc_type == "resume":
        _build_resume_docx(doc, lines)
    else:
        _build_cover_letter_docx(doc, lines, req.job_title, req.company)
    
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    
    safe_company = (req.company or "").replace(" ", "_")[:20]
    safe_title = (req.job_title or "").replace(" ", "_")[:20]
    
    if req.doc_type == "resume":
        filename = f"Resume_{safe_company}_{safe_title}.docx" if safe_company else "Optimized_Resume.docx"
    else:
        filename = f"Cover_Letter_{safe_company}_{safe_title}.docx" if safe_company else "Cover_Letter.docx"
    
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )


def _build_resume_docx(doc, lines):
    """Build a formatted resume .docx preserving AI optimization structure."""
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    
    # Section header patterns
    section_headers = [
        "summary", "objective", "experience", "work experience", "professional experience",
        "education", "skills", "technical skills", "certifications", "projects",
        "achievements", "awards", "publications", "volunteer", "interests",
        "professional summary", "core competencies", "qualifications",
    ]
    
    for i, line in enumerate(lines):
        stripped = line.strip()
        if not stripped:
            # Add minimal spacing for blank lines
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            continue
        
        lower = stripped.lower().rstrip(':').strip('─═-').strip()
        
        # Check if this is a name (first non-empty line, likely the candidate name)
        if i <= 2 and not any(c in stripped for c in ['•', '-', '|', '@', '●']) and len(stripped.split()) <= 5 and stripped == stripped.upper() or (i == 0 and len(stripped) < 40):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(stripped)
            run.bold = True
            run.font.size = Pt(16)
            run.font.color.rgb = RGBColor(30, 41, 59)
            p.paragraph_format.space_after = Pt(2)
            continue
        
        # Contact info line (contains email, phone, or | separators)
        if i <= 4 and ('|' in stripped or '@' in stripped or any(c.isdigit() for c in stripped[:3])):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(stripped)
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(100, 100, 100)
            p.paragraph_format.space_after = Pt(4)
            continue
        
        # Section headers
        if lower in section_headers or (stripped.endswith(':') and len(stripped.split()) <= 4) or stripped == stripped.upper() and len(stripped.split()) <= 4 and len(stripped) > 3:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(stripped.rstrip(':').upper())
            run.bold = True
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(55, 65, 81)
            # Add bottom border effect with a thin line
            p_after = doc.add_paragraph()
            p_after.paragraph_format.space_before = Pt(0)
            p_after.paragraph_format.space_after = Pt(3)
            run2 = p_after.add_run('─' * 70)
            run2.font.size = Pt(4)
            run2.font.color.rgb = RGBColor(200, 200, 200)
            continue
        
        # Bullet points
        if stripped.startswith(('•', '-', '●', '▪', '∙', '*')):
            bullet_text = stripped.lstrip('•-●▪∙* ').strip()
            p = doc.add_paragraph(style='List Bullet')
            run = p.add_run(bullet_text)
            run.font.size = Pt(10)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.left_indent = Pt(18)
            continue
        
        # Job title / Company lines (bold text with dates)
        if any(sep in stripped for sep in [' | ', ' – ', ' — ']) or (stripped.endswith(')') and '(' in stripped and any(month in stripped for month in ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec', 'Present'])):
            p = doc.add_paragraph()
            run = p.add_run(stripped)
            run.bold = True
            run.font.size = Pt(10.5)
            run.font.color.rgb = RGBColor(30, 41, 59)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(1)
            continue
        
        # Regular text
        p = doc.add_paragraph()
        run = p.add_run(stripped)
        run.font.size = Pt(10.5)
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)


def _build_cover_letter_docx(doc, lines, job_title="", company=""):
    """Build a formatted cover letter .docx."""
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    
    # Add date
    from datetime import datetime
    p = doc.add_paragraph()
    run = p.add_run(datetime.now().strftime("%B %d, %Y"))
    run.font.size = Pt(11)
    p.paragraph_format.space_after = Pt(12)
    
    is_first_para = True
    
    for line in lines:
        stripped = line.strip()
        
        if not stripped:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            is_first_para = False
            continue
        
        # Salutation (Dear ...)
        if stripped.lower().startswith('dear '):
            p = doc.add_paragraph()
            run = p.add_run(stripped)
            run.font.size = Pt(11)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(8)
            is_first_para = False
            continue
        
        # Closing (Sincerely, Best regards, etc.)
        if any(stripped.lower().startswith(c) for c in ['sincerely', 'best regards', 'regards', 'warm regards', 'respectfully', 'thank you']):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            run = p.add_run(stripped)
            run.font.size = Pt(11)
            p.paragraph_format.space_after = Pt(4)
            continue
        
        # Regular paragraph
        p = doc.add_paragraph()
        run = p.add_run(stripped)
        run.font.size = Pt(11)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = Pt(15)


@router.post("/ai/interview-prep")
async def get_interview_prep(request: Request, req: InterviewPrepRequest):
    """Generate interview preparation materials."""
    user = await get_current_user(request)
    
    # Check usage limits
    usage_check = await check_usage_limit(user.user_id, "interview_prep")
    if not usage_check["allowed"]:
        raise HTTPException(status_code=402, detail={
            "error": "usage_limit_reached",
            "feature": "interview_prep",
            "current": usage_check["current"],
            "limit": usage_check["limit"],
            "message": f"You've used your {usage_check['limit']} interview prep session this month. Upgrade to Pro for unlimited access."
        })
    
    profile = await db.user_profiles.find_one(
        {"user_id": user.user_id},
        {"_id": 0}
    )
    
    from emergentintegrations.llm.chat import LlmChat, UserMessage
    
    chat = LlmChat(
        api_key=EMERGENT_LLM_KEY,
        session_id=f"interview_{user.user_id}_{uuid.uuid4().hex[:8]}",
        system_message="""You are an expert career coach and interview preparation specialist.
Provide comprehensive interview preparation including common questions, 
STAR method examples, company research tips, and confidence-building advice."""
    ).with_model("openai", "gpt-5.2")
    
    # Handle v2 skills format (list of dicts with name, years, level)
    raw_skills = profile.get("skills", []) if profile else []
    skill_names = get_skill_names(raw_skills)
    skills = ", ".join(skill_names) if skill_names else "Not specified"
    
    prompt = f"""You are an expert interview coach creating a professional interview preparation document for a FAANG / enterprise role.

POSITION: {req.job_title} at {req.company}
JOB DESCRIPTION:
{req.job_description}
CANDIDATE SKILLS: {skills}

FORMATTING RULES (must follow exactly):

1. Use clean, professional Markdown
2. Use numbered, ALL-CAPS section headers (e.g., 1. COMMON INTERVIEW QUESTIONS)
3. Separate major sections with a horizontal rule (---)
4. Format each interview question using this exact structure:
   - Question as a level-4 header (####)
   - Suggested approach on one line (NO asterisks, NO italics)
   - Sample answer as a blockquote (>)
5. Keep sample answers concise: 4–6 sentences max
6. Use clear whitespace between questions
7. Do NOT use asterisks or italics
8. Do NOT use emojis or casual language
9. Make everything bold and easy to read

REQUIRED SECTIONS (in this exact order):

1. COMMON INTERVIEW QUESTIONS
- 5 common questions every interviewer asks
- Each with: #### Question, Suggested approach, > Sample answer

---

2. BEHAVIORAL QUESTIONS
- 5 behavioral questions using STAR method
- Each with: #### Question, STAR framework guidance, > Sample answer

---

3. TECHNICAL QUESTIONS
- 5 technical questions specific to {req.job_title}
- Each with: #### Question, Approach guidance, > Sample answer

---

4. INTERVIEW TIPS
- 5-7 tactical tips (numbered list)
- Focus on: preparation, body language, follow-up, negotiation

---

5. QUESTIONS TO ASK THE INTERVIEWER
- 5 intelligent questions (numbered list)
- Categories: role scope, team dynamics, growth, company direction

ANSWER STYLE:
- Professional, direct, data-driven
- Emphasize measurable impact and collaboration
- Use concrete examples
- Avoid generic phrases

Generate interview prep for {req.job_title} at {req.company} following this structure exactly."""
    
    try:
        response = await chat.send_message(UserMessage(text=prompt))
        await increment_usage(user.user_id, "interview_prep")
        return {"prep_materials": response}
    except Exception as e:
        logger.error(f"Interview prep error: {str(e)}")
        raise HTTPException(status_code=500, detail="Failed to generate interview prep")



# ========================
# CAREER PATH ANALYSIS
# ========================

@router.post("/ai/career-paths")
async def analyze_career_paths(request: Request):
    """Analyze user's resume and profile to suggest realistic career paths."""
    user = await get_current_user(request)

    # Check usage limits
    usage_check = await check_usage_limit(user.user_id, "career_paths")
    if not usage_check["allowed"]:
        raise HTTPException(status_code=402, detail={
            "error": "usage_limit_reached",
            "feature": "career_paths",
            "current": usage_check["current"],
            "limit": usage_check["limit"],
            "message": f"You've used your {usage_check['limit']} career path analysis this month. Upgrade to Pro for unlimited access."
        })

    profile = await db.user_profiles.find_one(
        {"user_id": user.user_id},
        {"_id": 0}
    )

    if not profile:
        raise HTTPException(status_code=400, detail="Please complete your profile first")

    profile = decrypt_sensitive_data(profile)

    # Check for cached analysis (within 7 days)
    cached = await db.career_analyses.find_one(
        {"user_id": user.user_id},
        {"_id": 0}
    )

    if cached:
        try:
            cached_date = datetime.fromisoformat(cached.get("created_at", ""))
            age_days = (datetime.now(timezone.utc) - cached_date).days
            if age_days < 7:
                logger.info(f"Returning cached career analysis (age: {age_days} days)")
                return cached.get("analysis")
        except Exception:
            pass

    resume_text = profile.get("resume_text", "")
    if not resume_text or len(resume_text) < 100:
        raise HTTPException(status_code=400, detail="Please upload your resume for career path analysis")

    skills = get_skill_names(profile.get("skills", []))
    experience_years = profile.get("experience_years", 0)
    current_titles = profile.get("job_titles", [])
    education = profile.get("highest_education", "")
    seniority = profile.get("seniority_level", "")
    industries = profile.get("industries", [])
    location = profile.get("address_city", "") or profile.get("address_state", "") or profile.get("address_country", "")

    from emergentintegrations.llm.chat import LlmChat, UserMessage

    chat = LlmChat(
        api_key=EMERGENT_LLM_KEY,
        session_id=f"career_paths_{user.user_id}_{uuid.uuid4().hex[:8]}",
        system_message="""You are an expert career advisor with deep knowledge of career transitions, salary data, and skill requirements.

Analyze a person's resume and suggest realistic career paths.

RULES:
1. Only suggest paths where they have 60%+ of required skills
2. Include 3-5 paths: current path advancement, adjacent moves, and 1-2 stretch roles
3. Base salary estimates on real market data (conservative, not optimistic)
4. Provide specific, actionable skill gaps

OUTPUT FORMAT (JSON only, no markdown):
{
  "current_path": {
    "title": "Business Analyst",
    "salary_avg": 75000,
    "market_demand": "high",
    "job_security": "stable",
    "growth_potential": "moderate"
  },
  "recommended_paths": [
    {
      "title": "Data Analyst",
      "match_score": 85,
      "category": "adjacent",
      "skills_you_have": ["SQL", "Excel", "Python"],
      "skills_to_learn": ["Tableau", "R", "Statistical Analysis"],
      "salary_range": {"min": 70000, "max": 95000, "avg": 85000},
      "salary_increase": "+$10k",
      "time_to_transition": "3-6 months",
      "market_demand": "high",
      "job_count_estimate": 1200,
      "difficulty": "moderate",
      "reasoning": "Your SQL and Python skills transfer directly. Adding Tableau would make you highly competitive.",
      "next_steps": [
        "Learn Tableau (free course, 20 hours)",
        "Build 2-3 portfolio projects with data visualization",
        "Apply to junior Data Analyst roles"
      ]
    }
  ]
}

CATEGORIES: "current" (advance in role), "adjacent" (easy transition), "stretch" (requires more effort)
MARKET DEMAND: "very_high", "high", "medium", "low"
TIME TO TRANSITION: "1-3 months", "3-6 months", "6-12 months", "1-2 years"
DIFFICULTY: "easy" (80%+ match), "moderate" (60-80%), "hard" (<60%)"""
    ).with_model("openai", "gpt-4o")

    skills_text = ", ".join(skills[:30]) if skills else "Not specified"
    titles_text = ", ".join(current_titles) if current_titles else "Not specified"
    industries_text = ", ".join(industries) if industries else "Any"

    prompt = f"""Analyze this person's career and suggest realistic next career paths.

PROFILE:
- Years of Experience: {experience_years}
- Current/Target Titles: {titles_text}
- Skills: {skills_text}
- Education: {education or 'Not specified'}
- Seniority: {seniority or 'Not specified'}
- Industries: {industries_text}
- Location: {location or 'Canada'}

RESUME (trimmed):
{resume_text[:4000]}

Suggest 3-5 realistic career paths with salary ranges for {location or 'Canadian market'}.
Return ONLY valid JSON, no markdown."""

    try:
        response = await chat.send_message(UserMessage(text=prompt))

        clean = response.strip()
        clean = re.sub(r'^```json\s*', '', clean)
        clean = re.sub(r'\s*```$', '', clean)
        import json as json_mod
        result = json_mod.loads(clean.strip())

        if "recommended_paths" not in result:
            raise ValueError("Missing recommended_paths")

        result["generated_at"] = datetime.now(timezone.utc).isoformat()
        result["user_location"] = location
        result["experience_years"] = experience_years

        await db.career_analyses.update_one(
            {"user_id": user.user_id},
            {"$set": {
                "user_id": user.user_id,
                "analysis": result,
                "created_at": datetime.now(timezone.utc).isoformat(),
                "resume_hash": hash(resume_text[:500]),
            }},
            upsert=True,
        )

        logger.info(f"Generated career path analysis for user {user.user_id}")
        await increment_usage(user.user_id, "career_paths")
        return result

    except ValueError as e:
        logger.error(f"Career paths parse error: {e}")
        raise HTTPException(status_code=500, detail="Failed to parse career analysis. Please try again.")
    except Exception as e:
        logger.error(f"Career path analysis error: {e}")
        raise HTTPException(status_code=500, detail=f"Analysis failed: {str(e)}")


@router.delete("/ai/career-paths")
async def clear_career_analysis(request: Request):
    """Clear cached career analysis to force regeneration."""
    user = await get_current_user(request)
    await db.career_analyses.delete_many({"user_id": user.user_id})
    return {"message": "Career analysis cache cleared"}



# ========================
# LEARNING RESOURCES
# ========================

LEARNING_RESOURCES = {
    # --- Data & Analytics ---
    "Tableau": [
        {"title": "Tableau Fundamentals", "provider": "Tableau Learning", "duration_hours": 20, "cost": "Free", "level": "beginner", "url": "https://www.tableau.com/learn/training", "type": "video_course"},
        {"title": "Tableau for Data Science", "provider": "Udemy", "duration_hours": 10, "cost": "$15", "level": "intermediate", "url": "https://www.udemy.com/course/tableau10/", "type": "video_course"},
    ],
    "Python": [
        {"title": "Python for Everybody", "provider": "Coursera", "duration_hours": 40, "cost": "Free (audit)", "level": "beginner", "url": "https://www.coursera.org/specializations/python", "type": "video_course"},
        {"title": "Automate the Boring Stuff with Python", "provider": "Al Sweigart", "duration_hours": 30, "cost": "Free", "level": "beginner", "url": "https://automatetheboringstuff.com/", "type": "book"},
        {"title": "Python Data Science Handbook", "provider": "Jake VanderPlas", "duration_hours": 25, "cost": "Free", "level": "intermediate", "url": "https://jakevdp.github.io/PythonDataScienceHandbook/", "type": "book"},
    ],
    "SQL": [
        {"title": "Intro to SQL", "provider": "Khan Academy", "duration_hours": 15, "cost": "Free", "level": "beginner", "url": "https://www.khanacademy.org/computing/computer-programming/sql", "type": "interactive"},
        {"title": "SQL for Data Analysis", "provider": "Mode Analytics", "duration_hours": 12, "cost": "Free", "level": "intermediate", "url": "https://mode.com/sql-tutorial/", "type": "interactive"},
    ],
    "R": [
        {"title": "R Programming", "provider": "Coursera", "duration_hours": 30, "cost": "Free (audit)", "level": "beginner", "url": "https://www.coursera.org/learn/r-programming", "type": "video_course"},
    ],
    "Power BI": [
        {"title": "Power BI Guided Learning", "provider": "Microsoft", "duration_hours": 25, "cost": "Free", "level": "beginner", "url": "https://learn.microsoft.com/en-us/power-bi/", "type": "interactive"},
    ],
    "Statistical Analysis": [
        {"title": "Statistics and Probability", "provider": "Khan Academy", "duration_hours": 30, "cost": "Free", "level": "beginner", "url": "https://www.khanacademy.org/math/statistics-probability", "type": "interactive"},
        {"title": "Statistics with R", "provider": "Coursera", "duration_hours": 20, "cost": "Free (audit)", "level": "intermediate", "url": "https://www.coursera.org/specializations/statistics", "type": "video_course"},
    ],
    "Data Visualization": [
        {"title": "Data Visualization with Python", "provider": "Coursera", "duration_hours": 18, "cost": "Free (audit)", "level": "intermediate", "url": "https://www.coursera.org/learn/python-for-data-visualization", "type": "video_course"},
    ],
    "Excel": [
        {"title": "Excel Skills for Business", "provider": "Coursera", "duration_hours": 30, "cost": "Free (audit)", "level": "intermediate", "url": "https://www.coursera.org/specializations/excel", "type": "video_course"},
        {"title": "Excel Formulas & Functions", "provider": "Microsoft", "duration_hours": 10, "cost": "Free", "level": "beginner", "url": "https://support.microsoft.com/en-us/excel", "type": "interactive"},
    ],
    "Advanced Data Analytics": [
        {"title": "Google Advanced Data Analytics Certificate", "provider": "Coursera", "duration_hours": 80, "cost": "Free (audit)", "level": "intermediate", "url": "https://www.coursera.org/professional-certificates/google-advanced-data-analytics", "type": "certificate"},
    ],
    # --- AI & Machine Learning ---
    "Machine Learning": [
        {"title": "Machine Learning", "provider": "Coursera (Andrew Ng)", "duration_hours": 60, "cost": "Free (audit)", "level": "intermediate", "url": "https://www.coursera.org/learn/machine-learning", "type": "video_course"},
        {"title": "fast.ai Practical Deep Learning", "provider": "fast.ai", "duration_hours": 40, "cost": "Free", "level": "intermediate", "url": "https://course.fast.ai/", "type": "video_course"},
    ],
    "Deep Learning": [
        {"title": "Deep Learning Specialization", "provider": "Coursera (Andrew Ng)", "duration_hours": 80, "cost": "Free (audit)", "level": "advanced", "url": "https://www.coursera.org/specializations/deep-learning", "type": "video_course"},
    ],
    "TensorFlow": [
        {"title": "TensorFlow Developer Certificate", "provider": "Coursera", "duration_hours": 60, "cost": "Free (audit)", "level": "intermediate", "url": "https://www.coursera.org/professional-certificates/tensorflow-in-practice", "type": "certificate"},
    ],
    "NLP": [
        {"title": "Natural Language Processing Specialization", "provider": "Coursera", "duration_hours": 50, "cost": "Free (audit)", "level": "advanced", "url": "https://www.coursera.org/specializations/natural-language-processing", "type": "video_course"},
    ],
    # --- Product & Management ---
    "Product Management": [
        {"title": "Product Management Fundamentals", "provider": "Product School", "duration_hours": 20, "cost": "Free", "level": "beginner", "url": "https://productschool.com/free-product-management-resources", "type": "video_course"},
        {"title": "Digital Product Management", "provider": "Coursera", "duration_hours": 15, "cost": "Free (audit)", "level": "intermediate", "url": "https://www.coursera.org/learn/uva-darden-digital-product-management", "type": "video_course"},
    ],
    "Product Lifecycle Management": [
        {"title": "Product Lifecycle Management Essentials", "provider": "LinkedIn Learning", "duration_hours": 8, "cost": "Free trial", "level": "intermediate", "url": "https://www.linkedin.com/learning/", "type": "video_course"},
    ],
    "Go-to-Market Strategies": [
        {"title": "Go-to-Market Strategy", "provider": "HubSpot Academy", "duration_hours": 6, "cost": "Free", "level": "intermediate", "url": "https://academy.hubspot.com/", "type": "video_course"},
    ],
    "Agile": [
        {"title": "Agile with Atlassian Jira", "provider": "Coursera", "duration_hours": 15, "cost": "Free (audit)", "level": "beginner", "url": "https://www.coursera.org/learn/agile-atlassian-jira", "type": "video_course"},
    ],
    "Project Management": [
        {"title": "Google Project Management Certificate", "provider": "Coursera", "duration_hours": 120, "cost": "$49/month", "level": "beginner", "url": "https://www.coursera.org/professional-certificates/google-project-management", "type": "certificate"},
        {"title": "Introduction to Project Management", "provider": "Coursera", "duration_hours": 18, "cost": "Free (audit)", "level": "beginner", "url": "https://www.coursera.org/learn/project-management-foundations", "type": "video_course"},
    ],
    "Scrum": [
        {"title": "Scrum Master Certification Prep", "provider": "Scrum.org", "duration_hours": 15, "cost": "Free", "level": "beginner", "url": "https://www.scrum.org/resources/scrum-guide", "type": "reading"},
    ],
    # --- Business & Finance ---
    "Financial Modeling": [
        {"title": "Financial Modeling Fundamentals", "provider": "Corporate Finance Institute", "duration_hours": 25, "cost": "Free", "level": "intermediate", "url": "https://corporatefinanceinstitute.com/resources/", "type": "video_course"},
    ],
    "Financial Analysis": [
        {"title": "Financial Analysis and Decision Making", "provider": "edX", "duration_hours": 30, "cost": "Free (audit)", "level": "intermediate", "url": "https://www.edx.org/learn/financial-analysis", "type": "video_course"},
    ],
    "Budgeting & Forecasting": [
        {"title": "Budgeting and Forecasting", "provider": "LinkedIn Learning", "duration_hours": 8, "cost": "Free trial", "level": "intermediate", "url": "https://www.linkedin.com/learning/", "type": "video_course"},
    ],
    "Accounting (GAAP/IFRS)": [
        {"title": "Introduction to Financial Accounting", "provider": "Coursera", "duration_hours": 25, "cost": "Free (audit)", "level": "beginner", "url": "https://www.coursera.org/learn/wharton-accounting", "type": "video_course"},
    ],
    "Business Analysis": [
        {"title": "Business Analysis Foundations", "provider": "LinkedIn Learning", "duration_hours": 10, "cost": "Free trial", "level": "beginner", "url": "https://www.linkedin.com/learning/", "type": "video_course"},
        {"title": "IIBA ECBA Study Guide", "provider": "IIBA", "duration_hours": 40, "cost": "Free", "level": "intermediate", "url": "https://www.iiba.org/business-analysis-certifications/ecba/", "type": "reading"},
    ],
    # --- Marketing ---
    "Digital Marketing": [
        {"title": "Google Digital Marketing Certificate", "provider": "Coursera", "duration_hours": 80, "cost": "Free (audit)", "level": "beginner", "url": "https://www.coursera.org/professional-certificates/google-digital-marketing-ecommerce", "type": "certificate"},
        {"title": "Digital Marketing Course", "provider": "HubSpot Academy", "duration_hours": 12, "cost": "Free", "level": "beginner", "url": "https://academy.hubspot.com/courses/digital-marketing", "type": "video_course"},
    ],
    "SEO/SEM": [
        {"title": "SEO Training Course", "provider": "HubSpot Academy", "duration_hours": 6, "cost": "Free", "level": "beginner", "url": "https://academy.hubspot.com/courses/seo-training", "type": "video_course"},
        {"title": "Google Ads Certification", "provider": "Google", "duration_hours": 15, "cost": "Free", "level": "intermediate", "url": "https://skillshop.withgoogle.com/", "type": "certificate"},
    ],
    "Content Marketing": [
        {"title": "Content Marketing Certification", "provider": "HubSpot Academy", "duration_hours": 8, "cost": "Free", "level": "beginner", "url": "https://academy.hubspot.com/courses/content-marketing", "type": "video_course"},
    ],
    "Google Analytics": [
        {"title": "Google Analytics Certification", "provider": "Google", "duration_hours": 15, "cost": "Free", "level": "intermediate", "url": "https://skillshop.withgoogle.com/", "type": "certificate"},
    ],
    # --- Technical / Engineering ---
    "JavaScript": [
        {"title": "JavaScript Algorithms & Data Structures", "provider": "freeCodeCamp", "duration_hours": 50, "cost": "Free", "level": "beginner", "url": "https://www.freecodecamp.org/learn/javascript-algorithms-and-data-structures/", "type": "interactive"},
        {"title": "The Odin Project - JavaScript", "provider": "The Odin Project", "duration_hours": 80, "cost": "Free", "level": "beginner", "url": "https://www.theodinproject.com/paths/full-stack-javascript", "type": "interactive"},
    ],
    "TypeScript": [
        {"title": "TypeScript Handbook", "provider": "Microsoft", "duration_hours": 15, "cost": "Free", "level": "intermediate", "url": "https://www.typescriptlang.org/docs/handbook/", "type": "reading"},
    ],
    "React": [
        {"title": "React Official Tutorial", "provider": "React.dev", "duration_hours": 20, "cost": "Free", "level": "beginner", "url": "https://react.dev/learn", "type": "interactive"},
    ],
    "Node.js": [
        {"title": "Node.js Tutorial", "provider": "freeCodeCamp", "duration_hours": 30, "cost": "Free", "level": "beginner", "url": "https://www.freecodecamp.org/learn/back-end-development-and-apis/", "type": "interactive"},
    ],
    "AWS": [
        {"title": "AWS Cloud Practitioner Essentials", "provider": "AWS", "duration_hours": 30, "cost": "Free", "level": "beginner", "url": "https://aws.amazon.com/training/", "type": "video_course"},
        {"title": "AWS Solutions Architect - Associate", "provider": "AWS", "duration_hours": 60, "cost": "Free", "level": "intermediate", "url": "https://aws.amazon.com/certification/certified-solutions-architect-associate/", "type": "certificate"},
    ],
    "Azure": [
        {"title": "Azure Fundamentals (AZ-900)", "provider": "Microsoft Learn", "duration_hours": 20, "cost": "Free", "level": "beginner", "url": "https://learn.microsoft.com/en-us/certifications/azure-fundamentals/", "type": "interactive"},
    ],
    "Docker": [
        {"title": "Docker Getting Started", "provider": "Docker", "duration_hours": 10, "cost": "Free", "level": "beginner", "url": "https://docs.docker.com/get-started/", "type": "interactive"},
    ],
    "Kubernetes": [
        {"title": "Kubernetes Basics", "provider": "Kubernetes.io", "duration_hours": 15, "cost": "Free", "level": "intermediate", "url": "https://kubernetes.io/docs/tutorials/kubernetes-basics/", "type": "interactive"},
    ],
    "Git/GitHub": [
        {"title": "Git & GitHub Crash Course", "provider": "freeCodeCamp", "duration_hours": 5, "cost": "Free", "level": "beginner", "url": "https://www.freecodecamp.org/news/git-and-github-crash-course/", "type": "video_course"},
    ],
    # --- Design ---
    "UI/UX Design": [
        {"title": "Google UX Design Certificate", "provider": "Coursera", "duration_hours": 100, "cost": "Free (audit)", "level": "beginner", "url": "https://www.coursera.org/professional-certificates/google-ux-design", "type": "certificate"},
    ],
    "Figma": [
        {"title": "Figma for Beginners", "provider": "Figma", "duration_hours": 10, "cost": "Free", "level": "beginner", "url": "https://help.figma.com/hc/en-us/categories/360002051613", "type": "interactive"},
    ],
    # --- Soft Skills & Leadership ---
    "Communication": [
        {"title": "Improving Communication Skills", "provider": "Coursera", "duration_hours": 12, "cost": "Free (audit)", "level": "beginner", "url": "https://www.coursera.org/learn/wharton-communication", "type": "video_course"},
    ],
    "Leadership": [
        {"title": "Foundations of Leadership", "provider": "edX", "duration_hours": 20, "cost": "Free (audit)", "level": "intermediate", "url": "https://www.edx.org/learn/leadership", "type": "video_course"},
    ],
    "Public Speaking": [
        {"title": "Introduction to Public Speaking", "provider": "Coursera", "duration_hours": 15, "cost": "Free (audit)", "level": "beginner", "url": "https://www.coursera.org/learn/public-speaking", "type": "video_course"},
    ],
    "Negotiation": [
        {"title": "Successful Negotiation", "provider": "Coursera", "duration_hours": 12, "cost": "Free (audit)", "level": "intermediate", "url": "https://www.coursera.org/learn/negotiation-skills", "type": "video_course"},
    ],
    # --- Healthcare ---
    "Medical Terminology": [
        {"title": "Medical Terminology Course", "provider": "Coursera", "duration_hours": 20, "cost": "Free (audit)", "level": "beginner", "url": "https://www.coursera.org/learn/medical-terminology", "type": "video_course"},
    ],
    # --- Cybersecurity ---
    "Cybersecurity": [
        {"title": "Google Cybersecurity Certificate", "provider": "Coursera", "duration_hours": 100, "cost": "Free (audit)", "level": "beginner", "url": "https://www.coursera.org/professional-certificates/google-cybersecurity", "type": "certificate"},
    ],
    # --- Supply Chain ---
    "Supply Chain Management": [
        {"title": "Supply Chain Management Specialization", "provider": "Coursera", "duration_hours": 50, "cost": "Free (audit)", "level": "intermediate", "url": "https://www.coursera.org/specializations/supply-chain-management", "type": "video_course"},
    ],
}


@router.get("/ai/learning-resources")
async def get_learning_resources(request: Request, skills: str = None):
    """Get curated learning resources for specific skill gaps.
    Query params:
      - skills: comma-separated skill names (e.g. "Tableau,Python,SQL")
    If omitted, returns the full catalogue of available skills.
    """
    await get_current_user(request)

    if not skills:
        return {
            "available_skills": sorted(LEARNING_RESOURCES.keys()),
            "total_resources": sum(len(r) for r in LEARNING_RESOURCES.values()),
        }

    skill_list = [s.strip() for s in skills.split(",") if s.strip()]
    resources_by_skill = {}

    for skill in skill_list:
        if skill in LEARNING_RESOURCES:
            resources_by_skill[skill] = LEARNING_RESOURCES[skill]
            continue
        for res_skill, res_list in LEARNING_RESOURCES.items():
            if skill.lower() in res_skill.lower() or res_skill.lower() in skill.lower():
                resources_by_skill[skill] = res_list
                break

    total_hours = 0
    total_free = 0
    total_paid = 0
    for resources in resources_by_skill.values():
        for r in resources:
            total_hours += r.get("duration_hours", 0)
            if "free" in r.get("cost", "").lower():
                total_free += 1
            else:
                total_paid += 1

    return {
        "skills_requested": skill_list,
        "resources": resources_by_skill,
        "summary": {
            "total_skills": len(resources_by_skill),
            "total_resources": total_free + total_paid,
            "total_hours": total_hours,
            "free_resources": total_free,
            "paid_resources": total_paid,
            "estimated_completion": f"{total_hours // 40} weeks at 40 hrs/week" if total_hours >= 40 else f"{total_hours} hours",
        },
    }


@router.get("/ai/career-paths/{path_title}/jobs")
async def get_jobs_for_career_path(
    request: Request,
    path_title: str,
    location: Optional[str] = None,
    min_match_score: int = 60,
):
    """
    Find real job listings for a career path and score user qualification.

    Path params:
      - path_title: career path title (e.g. "Data Analyst")
    Query params:
      - location: optional location filter (defaults to user profile city)
      - min_match_score: minimum match % to include (default 60)
    """
    user = await get_current_user(request)

    profile = await db.user_profiles.find_one(
        {"user_id": user.user_id}, {"_id": 0}
    )
    if not profile:
        raise HTTPException(status_code=400, detail="Profile required")

    profile = decrypt_sensitive_data(profile)

    if not location:
        location = (
            profile.get("address_city")
            or profile.get("address_state")
            or "Canada"
        )

    try:
        # --- 1. Search stored jobs (Greenhouse / Lever) ---
        title_words = [w for w in path_title.split() if len(w) >= 3]
        title_regex = "|".join(re.escape(w) for w in title_words) if title_words else re.escape(path_title)

        query_filter = {"title": {"$regex": title_regex, "$options": "i"}}
        if location:
            query_filter["$or"] = [
                {"location": {"$regex": location.split(",")[0].strip(), "$options": "i"}},
                {"is_remote": True},
            ]

        cached_jobs = await db.stored_jobs.find(
            query_filter, {"_id": 0}
        ).sort("posted_at", -1).limit(60).to_list(60)

        # --- 2. Search JSearch API for aggregator results ---
        jsearch_jobs = []
        if RAPIDAPI_KEY:
            try:
                async with httpx.AsyncClient(timeout=15.0) as client:
                    resp = await client.get(
                        "https://jsearch.p.rapidapi.com/search",
                        params={
                            "query": f"{path_title} {location}",
                            "num_pages": "2",
                        },
                        headers={
                            "X-RapidAPI-Key": RAPIDAPI_KEY,
                            "X-RapidAPI-Host": "jsearch.p.rapidapi.com",
                        },
                    )
                    if resp.status_code == 200:
                        jsearch_jobs = resp.json().get("data", [])[:30]
            except Exception:
                logger.warning("JSearch API unavailable for career-path job search")

        # --- 3. Normalise JSearch results into the same shape ---
        normalised_jsearch = []
        for j in jsearch_jobs:
            city = j.get("job_city") or ""
            state = j.get("job_state") or ""
            normalised_jsearch.append({
                "job_id": j.get("job_id"),
                "title": j.get("job_title"),
                "company": j.get("employer_name"),
                "location": f"{city}, {state}".strip(", "),
                "description": (j.get("job_description") or "")[:500],
                "apply_link": j.get("job_apply_link"),
                "posted_at": j.get("job_posted_at_datetime_utc"),
                "source": "aggregator",
            })

        all_jobs = cached_jobs + normalised_jsearch

        # --- 4. Deduplicate ---
        seen = set()
        unique_jobs = []
        for job in all_jobs:
            jid = job.get("job_id") or job.get("title", "") + job.get("company", "")
            if jid not in seen:
                seen.add(jid)
                unique_jobs.append(job)

        # --- 5. Score each job against user profile ---
        matched_jobs = []
        for job in unique_jobs[:60]:
            job_for_match = {
                "job_title": job.get("title"),
                "employer_name": job.get("company"),
                "job_description": job.get("description") or job.get("title", ""),
                "job_city": (job.get("location") or "").split(",")[0].strip(),
                "job_state": (job.get("location") or "").split(",")[-1].strip() if "," in (job.get("location") or "") else "",
                "job_is_remote": "remote" in (job.get("location") or "").lower(),
            }

            match_eval = evaluate_job_match(job_for_match, profile)

            if match_eval["score"] >= min_match_score:
                matched_jobs.append({
                    "job": {
                        "job_id": job.get("job_id"),
                        "title": job.get("title"),
                        "company": job.get("company"),
                        "location": job.get("location"),
                        "apply_link": job.get("apply_link") or job.get("url"),
                        "posted_at": job.get("posted_at"),
                        "source": job.get("source", "ats_board"),
                    },
                    "match_score": match_eval["score"],
                    "match_recommendation": match_eval["recommendation"],
                    "strengths": match_eval["strengths"],
                    "gaps": match_eval["gaps"],
                    "grounded_strengths": match_eval.get("grounded_strengths", []),
                    "matched_skills": match_eval.get("matched_skills", []),
                    "ready_to_apply": match_eval["score"] >= 75,
                })

        matched_jobs.sort(key=lambda x: x["match_score"], reverse=True)
        top_jobs = matched_jobs[:30]

        return {
            "career_path": path_title,
            "location": location,
            "total_jobs_found": len(unique_jobs),
            "qualified_jobs": len(matched_jobs),
            "jobs": top_jobs,
            "summary": {
                "ready_to_apply_now": sum(1 for j in top_jobs if j["match_score"] >= 75),
                "close_match": sum(1 for j in top_jobs if 65 <= j["match_score"] < 75),
                "stretch_roles": sum(1 for j in top_jobs if j["match_score"] < 65),
                "avg_match_score": round(
                    sum(j["match_score"] for j in top_jobs) / len(top_jobs), 1
                ) if top_jobs else 0,
            },
        }

    except Exception as e:
        logger.error(f"Career-path job search error: {e}")
        raise HTTPException(status_code=500, detail=f"Job search failed: {str(e)}")




# ========================
# JOB COMPARISON / ANALYZE MATCH ROUTES
# ========================

