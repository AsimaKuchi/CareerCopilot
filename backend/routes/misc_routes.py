"""
routes/misc_routes.py - Miscellaneous endpoints: root, health, test pages, analytics, admin user mgmt.

Endpoints (under /api):
  GET  /                                    -- Root info
  GET  /health                              -- Liveness probe
  GET  /extension/download                  -- Download bundled extension zip
  GET  /test-download                       -- Test file download (debug)
  GET  /test-form                           -- Sample form (autofill testing)
  POST /analytics/company-search            -- Track company filter usage
  GET  /admin/users                         -- Paginated user list (admin)
  PUT  /admin/users/{user_id}/role          -- Promote/demote user (admin)
  GET  /admin/stats                         -- Platform stats (admin)
  POST /admin/unlock-user                   -- Unlock a locked-out account (admin)

Refactored from monolithic server.py (Feb 2026).
"""
from datetime import datetime, timezone, timedelta
from typing import Optional
import os
import uuid

from fastapi import APIRouter, HTTPException, Request
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, RedirectResponse, Response

from core import db, logger, STATIC_DOWNLOADS_DIR, get_current_user, get_admin_user

router = APIRouter()


@router.get("/")
async def root():
    return {"message": "MyCareerCoPilot API", "status": "healthy"}

@router.get("/health")
async def health():
    return {"status": "healthy"}


@router.get("/extension/download")
async def download_extension():
    path = "/app/browser-extension.zip"
    if not os.path.exists(path):
        raise HTTPException(status_code=404, detail="Extension file not found")
    return FileResponse(path, filename="MyCareerCoPilot-Extension.zip", media_type="application/zip")


@router.get("/test-download")
async def test_download():
    """Test endpoint - downloads a simple DOCX file without authentication."""
    from docx import Document
    import io
    
    # Create a simple test document
    doc = Document()
    doc.add_heading('Test Download', 0)
    doc.add_paragraph('If you can read this, downloads are working!')
    doc.add_paragraph('Generated at: ' + datetime.now().isoformat())
    
    # Save to bytes
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    docx_bytes = buffer.getvalue()
    
    # Save to static downloads
    filename = f"test_download_{uuid.uuid4().hex[:8]}.docx"
    filepath = os.path.join(STATIC_DOWNLOADS_DIR, filename)
    with open(filepath, 'wb') as f:
        f.write(docx_bytes)
    
    # Redirect to static file
    return RedirectResponse(url=f"/api/static-downloads/{filename}", status_code=302)

# Test form for autofill bot testing
@router.get("/test-form")
async def serve_test_form():
    """Serve a test application form for bot testing."""
    from fastapi.responses import HTMLResponse
    html = """<!DOCTYPE html>
<html>
<head>
    <title>Test Application Form</title>
    <style>
        body { font-family: Arial; max-width: 600px; margin: 50px auto; padding: 20px; }
        .field { margin: 15px 0; }
        label { display: block; margin-bottom: 5px; font-weight: bold; }
        input, textarea, select { width: 100%; padding: 8px; box-sizing: border-box; }
        textarea { height: 100px; }
        button { background: #007bff; color: white; padding: 10px 20px; border: none; cursor: pointer; margin-top: 20px; }
    </style>
</head>
<body>
    <h1>Test Job Application Form</h1>
    <form id="application-form">
        <div class="field">
            <label for="first_name">First Name</label>
            <input type="text" id="first_name" name="first_name">
        </div>
        <div class="field">
            <label for="last_name">Last Name</label>
            <input type="text" id="last_name" name="last_name">
        </div>
        <div class="field">
            <label for="email">Email Address</label>
            <input type="email" id="email" name="email">
        </div>
        <div class="field">
            <label for="phone">Phone Number</label>
            <input type="tel" id="phone" name="phone">
        </div>
        <div class="field">
            <label for="city">City</label>
            <input type="text" id="city" name="city">
        </div>
        <div class="field">
            <label for="linkedin">LinkedIn URL</label>
            <input type="url" id="linkedin" name="linkedin" placeholder="https://linkedin.com/in/...">
        </div>
        <div class="field">
            <label for="github">GitHub URL</label>
            <input type="url" id="github" name="github" placeholder="https://github.com/...">
        </div>
        <div class="field">
            <label for="portfolio">Portfolio / Website</label>
            <input type="url" id="portfolio" name="portfolio" placeholder="https://...">
        </div>
        <div class="field">
            <label for="resume">Resume/CV</label>
            <input type="file" id="resume" name="resume" accept=".pdf,.docx">
        </div>
        <div class="field">
            <label for="cover_letter">Cover Letter</label>
            <textarea id="cover_letter" name="cover_letter" placeholder="Write your cover letter..."></textarea>
        </div>
        <button type="submit">Submit Application</button>
    </form>
</body>
</html>"""
    return HTMLResponse(content=html)


# ========================
# COMPANY SEARCH ANALYTICS
# ========================

@router.post("/analytics/company-search")
async def track_company_search(request: Request):
    """Track which companies users search for (helps prioritize job sourcing)."""
    user = await get_current_user(request)
    body = await request.json()
    company = body.get("company", "").strip()
    
    if not company:
        return {"status": "skipped"}
    
    await db.company_search_analytics.update_one(
        {"company": company.lower()},
        {
            "$inc": {"search_count": 1},
            "$set": {"last_searched": datetime.now(timezone.utc).isoformat()},
            "$addToSet": {"users": user.user_id},
        },
        upsert=True,
    )
    return {"status": "tracked"}


# ========================
# ADMIN ENDPOINTS
# ========================

@router.get("/admin/users")
async def admin_list_users(request: Request, page: int = 1, limit: int = 20):
    """List all users (admin only)."""
    admin = await get_admin_user(request)
    skip = (page - 1) * limit
    total = await db.users.count_documents({})
    users = await db.users.find(
        {},
        {"_id": 0, "password_hash": 0, "reset_token": 0, "verification_token": 0}
    ).sort("created_at", -1).skip(skip).limit(limit).to_list(limit)
    return {"users": users, "total": total, "page": page, "pages": (total + limit - 1) // limit}

@router.put("/admin/users/{user_id}/role")
async def admin_set_user_role(user_id: str, request: Request):
    """Set a user's role (admin only)."""
    admin = await get_admin_user(request)
    body = await request.json()
    new_role = body.get("role")
    if new_role not in ("user", "admin"):
        raise HTTPException(status_code=400, detail="Role must be 'user' or 'admin'")
    if admin.user_id == user_id and new_role != "admin":
        raise HTTPException(status_code=400, detail="Cannot remove your own admin role")
    result = await db.users.update_one({"user_id": user_id}, {"$set": {"role": new_role}})
    if result.matched_count == 0:
        raise HTTPException(status_code=404, detail="User not found")
    return {"message": f"User role updated to {new_role}"}

@router.get("/admin/stats")
async def admin_stats(request: Request):
    """Get admin dashboard statistics (admin only)."""
    await get_admin_user(request)
    total_users = await db.users.count_documents({})
    total_jobs = await db.stored_jobs.count_documents({})
    total_apps = await db.applications.count_documents({})
    total_sessions = await db.user_sessions.count_documents({})
    recent_signups = await db.users.count_documents({
        "created_at": {"$gte": (datetime.now(timezone.utc) - timedelta(days=7)).isoformat()}
    })
    return {
        "total_users": total_users,
        "total_jobs_indexed": total_jobs,
        "total_applications": total_apps,
        "active_sessions": total_sessions,
        "signups_last_7_days": recent_signups
    }

@router.post("/admin/unlock-user")
async def admin_unlock_user(request: Request):
    """Unlock a locked user account (admin only)."""
    await get_admin_user(request)
    body = await request.json()
    email = body.get("email")
    if not email:
        raise HTTPException(status_code=400, detail="Email required")
    result = await db.users.update_one(
        {"email": email.lower()},
        {"$set": {"failed_login_attempts": 0}, "$unset": {"locked_until": ""}}
    )
    if result.matched_count == 0:
        raise HTTPException(status_code=404, detail="User not found")
    return {"message": f"Account {email} unlocked"}

